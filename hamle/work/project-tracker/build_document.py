from pathlib import Path
import json
from decimal import Decimal, ROUND_HALF_UP
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE=Path(__file__).resolve().parent
ROOT=HERE.parent.parent
D=json.loads((HERE/'project-plan.json').read_text())
OUT=ROOT/'outputs/hamle-yeni-mac-devir-2026-10-10'
OUT.mkdir(parents=True,exist_ok=True)
doc=Document()
s=doc.sections[0]
s.page_width=Inches(8.5); s.page_height=Inches(11)
s.top_margin=Inches(.67); s.bottom_margin=Inches(.65)
s.left_margin=Inches(.7); s.right_margin=Inches(.7)
s.header_distance=Inches(.28); s.footer_distance=Inches(.28)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 st=doc.styles[name]; st.font.name='Arial'; st.font.color.rgb=RGBColor.from_string('000000')
st=doc.styles['Normal']; st.font.size=Pt(11); st.paragraph_format.line_spacing=1.09; st.paragraph_format.space_after=Pt(7)
for name,size in [('Title',27),('Heading 1',19),('Heading 2',13),('Heading 3',11.5)]:
 st=doc.styles[name]; st.font.size=Pt(size); st.font.bold=True
 st.paragraph_format.space_before=Pt(13 if name!='Title' else 0); st.paragraph_format.space_after=Pt(8)
 st.paragraph_format.keep_with_next=True
doc.styles['Subtitle'].font.size=Pt(13)
doc.styles['Subtitle'].font.italic=False
doc.styles['Subtitle'].paragraph_format.space_after=Pt(13)
for st in doc.styles:
 for border in list(st.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
doc.core_properties.title='Gamechangers Hamle Proje Takip ve Uygulama Planı'
doc.core_properties.author='Gamechangers'
doc.core_properties.subject='9 Ekim 2026 durum değerlendirmesi ve uygulama başlangıç planı'
head=s.header.paragraphs[0]; head.text='GAMECHANGERS   /   HAMLE PROJE TAKİBİ'
head.runs[0].font.size=Pt(8); head.runs[0].font.color.rgb=RGBColor.from_string('687587')
foot=s.footer.paragraphs[0]; foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
run=foot.add_run('Şirket içi • 10 Ekim 2026 • '); run.font.size=Pt(8)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); foot._p.append(fld)

def p(t='',bold=False,small=False):
 pp=doc.add_paragraph(); rr=pp.add_run(t); rr.bold=bold
 if small: rr.font.size=Pt(9)
 return pp
def h(t,level=1): return doc.add_heading(t,level=level)
def page(title,kicker=None):
 doc.add_page_break()
 if kicker: p(kicker,small=True)
 h(title)
def table(headers,rows,widths=None,size=10):
 t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
 widths=widths or [7.1/len(headers)]*len(headers)
 for c,w in zip(t.columns,widths): c.width=Inches(w)
 for c,w,txt in zip(t.rows[0].cells,widths,headers): c.width=Inches(w); c.text=txt
 for vals in rows:
  cc=t.add_row().cells
  for c,w,txt in zip(cc,widths,vals): c.width=Inches(w); c.text=str(txt)
 for ri,row in enumerate(t.rows):
  trPr=row._tr.get_or_add_trPr()
  cant=OxmlElement('w:cantSplit'); trPr.append(cant)
  if ri==0:
   repeat=OxmlElement('w:tblHeader'); trPr.append(repeat)
  for ci,c in enumerate(row.cells):
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr()
   sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'172A46' if ri==0 else ('F1F5FA' if ri%2==0 else 'FFFFFF')); pr.append(sh)
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
   pr.append(borders)
   mar=OxmlElement('w:tcMar')
   for edge,val in [('top',85),('bottom',85),('left',100),('right',100)]:
    e=OxmlElement('w:'+edge); e.set(qn('w:w'),str(val)); e.set(qn('w:type'),'dxa'); mar.append(e)
   pr.append(mar)
   for pp in c.paragraphs:
    pp.paragraph_format.space_after=Pt(0); pp.paragraph_format.line_spacing=1.03
    if ci==0 and widths[0]<1: pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for rr in pp.runs:
     rr.font.name='Arial'; rr.font.size=Pt(size); rr.bold=(ri==0)
     rr.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '172432')
 p('',small=True).paragraph_format.space_after=Pt(0)
 return t
def taskrows(ids):
 rows=[]
 for t in D['tasks']:
  if t['id'] not in ids: continue
  rows.append([t['id'],t['title'],('Tamamlandı' if t['done'] else t['nextAction'])])
 return rows
def phase_rows(prefix,ids=None):
 ts=ids or [x['id'] for x in D['tasks'] if x['phaseId']==prefix]
 table(['Kod','İş','Durum veya sıradaki somut adım'],taskrows(ts),[.48,2.66,3.96],9.6)
def link(label,url):
 pp=doc.add_paragraph(); hl=OxmlElement('w:hyperlink')
 rid=pp.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
 hl.set(qn('r:id'),rid); r=OxmlElement('w:r'); rPr=OxmlElement('w:rPr')
 c=OxmlElement('w:color'); c.set(qn('w:val'),'244E78'); rPr.append(c)
 sz=OxmlElement('w:sz'); sz.set(qn('w:val'),'18'); rPr.append(sz); r.append(rPr)
 txt=OxmlElement('w:t'); txt.text=label; r.append(txt); hl.append(r); pp._p.append(hl)
 pp.paragraph_format.space_after=Pt(5)

done=sum(t['done'] for t in D['tasks']); pts=sum(t['weight']*t['done'] for t in D['tasks'])
total=sum(t['weight'] for t in D['tasks']); count=len(D['tasks'])
remaining=count-done; progress=pts/total*100; remaining_points=total-pts
assert count==65 and total==100
assert done==D['currentSummary']['completedTasks'] and pts==D['currentSummary']['completedWeight']

doc.add_heading('Gamechangers Hamle',0)
doc.add_paragraph('Proje takip ve uygulama planı',style='Subtitle')
p('10 Ekim 2026 devir sürümü • İşlem kanıtları 9 Ekim 23.59 • Şirket içi',small=True)
h('Şu an neredeyiz',1)
p(f'Bu belgede kurulan 100 puanlık iş planında {pts:g} puan tamamlandı; {remaining_points:g} puan açık. Güncel ilerleme %{progress:g}. {count} işin {done} adedi kendi dar kapsamlı kabul ölçütüne göre bitti. Bulutta bağımsız çalışan ve canlı kabulü tamamlanmış ajan sayısı henüz 0 / 4.',bold=True)
p('Bu yüzde, yönetim için önerilen başlangıç ölçüsünün güncel durumudur. Geçen süreyi, öğrenme seviyesini veya canlıya hazırlığı ölçmez. Bağlantı, kalite ve yetki koşulları ayrıca geçilmeden müşteri operasyonu açılmaz.')
table(['Tamamlanan temel işler','Önümüzdeki temel işler'],[
 ['Rol ve yetki kuralları; kaynaklı ilk öğrenme; 20 yazılı mail provası.','Güncel marka dosyaları; gerçek model ve tasarım araçları; bağlantı iptal kontrolleri.'],
 ['Gmail API iç gönderim, teslim ve yenileme; Basecamp API görev, atama, dosya ve revizyon testi.','Bilgisayar kapalıyken çalışan bulut hizmeti; kayıt, zamanlama ve maliyet kontrolü.'],
 ['Cloud projesi, Gmail API ve OAuth uygulama kaydı; yerel testler.','Hamle ile uzmanların uçtan uca üretimi; gölge çalışma ve kapsam bazlı canlı kabul.']],[3.55,3.55],10.2)
h('İlk somut sıra',2)
p('1. Onaylı 50 USD sınırını uygulamaya taşı; Astra + max hesap erişimini ve gerçek çağrıyı doğrula.\n2. Marka kişilerini, Basecamp projelerini, dosyalarını ve yetki sınırlarını tamamla.\n3. Bulut ortamını, tam posta işleyicisini ve üretim araçlarını bağla.\n4. İç talebi metin, grafik, revizyon ve teslimata kadar çalıştır.\n5. Gölge sonuçları ve kritik kontroller geçerse rutin müşteri kapsamını aç.')
p('Pazartesi 12 Ekim, ilk üç rolün sınırları belli rutin işleri için koşullu hedef olabilir. Dört rolün tüm iş türlerinde %100 bağımsız çalışması için henüz doğrulanmış bir tarih yok.',bold=True)

page('Kanıtlanan durum ve açık alanlar')
p('Bu sürüm, 9 Ekim gerçek bağlantı testlerini ve yeni kullanıcı kararlarını takip planına işler. İfadeler kayıtlı test anlarını anlatır. En yeni özel kanıt, eski genel özetin önüne alınır; belge yenilemesi kendi başına hesap veya müşteri işlemi yapmaz.')
table(['Alan','Kanıtlanan','Henüz kanıtlanmayan'],[
 ['Gmail','Hamle OAuth; iç gönderim ve Ahmet’te teslim; yenileme ve ayrı süreçte okuma geçti. Kendi gönderdiği mesaj engellendi.','Gerçek izin iptali; gelen posta ve eklerin uçtan uca işlenmesi; otomatik posta akışı.'],
 ['Basecamp','İç projede API görev, Hamle’ye atama, R1 revizyon, dosya yükleme ve tam içerik geri okuma geçti.','Erişim yenileme ve iptal; marka projeleri; üretim iş akışı.'],
 ['Google Cloud','gamechangers-hamle projesi; Gmail API; OAuth uygulaması ve Desktop istemcisi.','Çalışır bulut uygulaması; zamanlayıcı; veri ve sır hizmetleri.'],
 ['Bilgi hazırlığı','105 benzersiz yazışma dizisi ve 763 mesaj kaydı; ek 60 ayrıntılı konu.','Tüm arşiv kapsamı; görsel eklerin incelenmesi; tüm markaların güncel ve onaylı dosyaları.'],
 ['Kurallar','35 operasyon, 30 Gmail yardımcı ve 54 posta normalleştirme/okuma yerel testi geçti.','Gerçek uçtan uca müşteri akışı, model karar kalitesi ve üretim kalitesi.'],
 ['Ajanlar','Dört rol tanımı ve yazılı örnekler hazır.','Kalıcı bulut üzerinde kabul edilmiş bağımsız dört ajan.']],[1.02,3.0,3.08],10.1)
h('Öğrenme kapsamının doğru okunması',2)
p('Ek 60 ayrıntılı konu: WAT 27, Rabam 26, eski ya da/Multicosmetics/Remedium grubunda 7. Bu yedi konunun yeni marka ayrımına kaynak bazında eşlenmesi açık. Kayıtlı mesaj aralığı 12 Mayıs–7 Ekim 2026; tam bir yıl veya tüm arşiv tamamlanmış sayılmaz. Öğrenmeye yüzde vermiyoruz.')
p('Bir DOCX’in çıkarılmış metni ve iki bağlı belgenin metni okunmuş; görsel olarak incelenmiş ek sayısı kayıtlı kapsamda sıfır. Dosya adını görmek veya metni çıkarmak, görsel ve videoyu incelemek yerine geçmez.')
p('Güncellik notu: 8 Ekim kurulum özeti bazı bağlantılarda eskidir. Son Gmail kanıtları geçerlidir. Tam dizi okumasında yalnız kendi giden mesajının engellenmesi sınandı; gerçek gelen posta, ekler ve görünmeyen CC değişimi açık kaldı.',small=True)

page('İlerleme hesabı ve iş paketleri')
p('Ağırlıklar bu takip başlangıcı için önerilmiştir. İşin bağımlılığı, uygulama yükü ve kabul riski dikkate alınarak toplam 100 puan dağıtıldı. Bir iş ancak kendi kabul ölçütü ve kanıtı varsa 1; diğer durumlarda 0 alır. Devam eden işlere tahmini kısmi puan verilmedi.')
rows=[]
for ph in D['phases']:
 ts=[t for t in D['tasks'] if t['phaseId']==ph['id']]; phase_total=sum(t['weight'] for t in ts); got=sum(t['weight']*t['done'] for t in ts)
 percent=(Decimal(str(got))*100/Decimal(str(phase_total))).quantize(Decimal('0.1'),rounding=ROUND_HALF_UP)
 rows.append([ph['id'],ph['name'],f'{phase_total:g}',f'{got:g}',('%'+str(percent)).replace('.',','),f'{sum(t["done"] for t in ts)} / {len(ts)}'])
rows.append(['','Toplam',f'{total:g}',f'{pts:g}',f'%{progress:g}',f'{done} / {count}'])
table(['Kod','İş paketi','Puan','Biten puan','İlerleme','Biten iş'],rows,[.42,3.12,.62,.84,.87,1.23],10.2)
p(f'Hesap: tamamlanan işlerin ağırlık toplamı / tüm işlerin ağırlık toplamı = {pts:g} / {total:g} = %{progress:g}. Kalan {remaining} iş toplam {remaining_points:g} puandır. Basit iş sayısı oranı {done} / {count} ≈ %{done/count*100:.1f}; ana gösterge ağırlıklı orandır.')
p('V1 tarihsel başlangıç: 15 / 65 iş ve 22 / 100 puan. İlk dosyalar değiştirilmedi; bu sürümde yeni kabul kanıtları ve kararlar işlendi.',small=True)
h('Ölçünün sınırları',2)
p('Cloud projesi açma ile bulut işletimini bitirme aynı ağırlıkta değildir. Rol tanımı hazırlığı yalnız A paketinde sayılır; F, G ve H paketlerinde o rollerin gerçek üretimi aranır. Okuma ve kaynaklı sentez farklı teslimlerdir; hiçbirisi tüm arşivin öğrenildiğini göstermez.')
p('Bu belge tam dört rollü vizyonu takip eder. İlk Hamle + metin + basit grafik açılışı daha dar bir kapsamdır; o kapsamın kabulü toplam projeyi %100 yapmaz. Sonradan kapsam eklenirse payda ve sürüm birlikte değiştirilir.')

page('Sistem nasıl çalışacak')
p('Google Cloud, bu işin bilgisayar kapalıyken de çalışabileceği işletim ortamıdır. Ajan ise rol talimatları, marka bilgisi, yapay zekâ modeli, kullanabildiği araçlar ve kontrol kurallarının birleşimidir. Bu konuşma ve tarayıcı oturumu kendi başına kalıcı hizmet oluşturmaz.')
table(['Katman','Yapacağı iş','Kurulum yaklaşımı'],[
 ['Gmail','Hamle’nin katıldığı talepleri almak; doğru zincirde yanıtlamak.','Yalnız Hamle canlı posta kimliği. Ahmet geçmiş öğrenme kaynağı olarak ayrı tutulur.'],
 ['Hamle','Yetkiyi kontrol etmek; talebi anlamak; brief, takip, devir ve iletişim.','Her işte marka, kaynak mesaj, kapsam, sahip, süre ve durum kaydı.'],
 ['Uzmanlar','Metin, grafik ve animasyon üretmek; kendi kalite kontrolünü yapmak.','Her role yalnız gerekli marka bilgisi ve üretim aracı verilir.'],
 ['Basecamp','İşin sahibi, son sürümü, revizyonu, kararı ve durumunu göstermek.','Tek iş kaydı ve tek güncel sürüm; iç ve müşteri görünürlüğü ayrı.'],
 ['Bulut işletimi','Zamanında başlatmak, hatayı yakalamak, güvenli devam etmek.','Cloud Run + zamanlayıcı; kalıcı iş kaydı, dosya deposu ve sır kasası önerisi.']],[1.05,2.95,3.1],10.2)
h('Bir talebin adımları',2)
p('1. Gerçek To/CC ve işlem yetkisi kontrol edilir.\n2. Marka, iş türü, risk ve eksik bilgi belirlenir.\n3. Talep tek bir iş kaydına bağlanır; açık brief hazırlanır.\n4. Basecamp’te görev açılır veya mevcut iş güncellenir.\n5. Uzman üretir; çıktı kendi kalite kontrolünden geçer.\n6. Hamle son sürümü, alıcıyı, kapsamı ve güncel katılımı yeniden kontrol eder.\n7. Gerekli müşteri değerlendirmesi ve revizyonlar aynı kayıtta yürür.\n8. Teslim, onay, kapanış ve öğrenme kaydı tamamlanır.')
p('Rutin işler, tanımlı yetki içinde senden her seferinde onay beklemez. Belirsiz kapsam, fiyat, şikâyet, strateji ve durdurma olayları ayrı yollara gider. Bulut bileşenleri henüz kurulmuş değil; bu bölüm hedef tasarımdır.',small=True)

page('Dört ajanı nasıl hazırlayacağız')
p('Buradaki eğitim, geçmiş mailleri bir modele yükleyip her şeyi bildiğini varsaymak değildir. Kaynaklı marka kayıtları, açık görev tanımları, iyi örnekler, araç yetkileri ve ölçülebilir denemeler birlikte hazırlanır. İlk kurulumda özel model eğitimi zorunlu kabul edilmez.')
table(['Rol','Girdi ve çıktı','Kendi kontrolü'],[
 ['Hamle','Mail ve iş geçmişi → sınıflandırma, brief, görev, yanıt, takip.','Güncel To/CC, marka, yetkili kişi, yönlendirme, alıcı, sürüm, teslim kaydı.'],
 ['Metin yazarı','Onaylı brief ve marka kaynakları → başlık, gövde, alternatif, revizyon.','İddianın kaynağı, ton, yazım, platform sınırı, brief uyumu.'],
 ['Grafik tasarımcı','Onaylı metin, logo ve varlıklar → görsel, ölçü varyantı, düzenlenebilir dosya.','Logo, Türkçe metin, ölçü, okunurluk, lisans, dışa aktarma.'],
 ['Animasyon tasarımcı','Onaylı metin ve görsel dil → storyboard, kurgu, sesli/sessiz çıktı.','Süre, ritim, ses hakları, altyazı, logo, kareler ve son dosya.']],[1.18,3.04,2.88],10.2)
h('Hazırlık döngüsü',2)
p('1. Her marka için onaylı kılavuz, örnek ve güncel hizmet kapsamı toplanır.\n2. Bilgiler kaynak tarihi ve onay durumuyla kaydedilir; çelişkiler açık bırakılır.\n3. Marka bilgisi, iş kayıtları, rol kuralları ve araç izinleri ayrı tutulur.\n4. Önce normal örnekler, sonra belirsiz ve riskli talepler denenir.\n5. Daha önce cevabı gösterilmemiş eski taleplerde kör karşılaştırma yapılır.\n6. Gerçek iç işte üretim ve revizyon denenir; hataya göre kayıt veya kural düzeltilir.\n7. Canlıda yalnız kabul edilen marka ve iş türü açılır; yeni sürüm tekrar sınanır.')
p('Cansu iletişim, Özge sosyal plan, atanacak tasarım sorumlusu görsel kalite için önerilen değerlendirme sahipleridir. Kabul sorumlulukları teyit edilmelidir. Resmî katalogda GPT 6 Astra bulundu; Ultra listelenmediği için kullanıcı Astra + max denemesini onayladı. Hesap erişimi ve gerçek API çağrısı henüz doğrulanmadı.')

page('Yetki sınırları ve katılım kuralı')
table(['Talep türü','İşleyiş'],[
 ['Teklif, fiyat, indirim, ek kapsam veya ticari taahhüt','Ahmet’e yönlendirilir. Hamle kendi başına fiyat veya teslim taahhüdü vermez.'],
 ['Sorun, şikâyet, kriz veya stratejik karar','Cansu ve Ahmet’e yönlendirilir; sorun, etkisi ve önerilen çözüm özetlenir.'],
 ['Sosyal gündem, içerik planı, trend ve dönemsel öneri','Özge ile birlikte yürütülür. Özge’nin görüşü gerçekten alınmadan alınmış gibi yazılmaz.'],
 ['Yetkili, tanımlı rutin iş','Hamle doğrudan yürütür; uzmanlara görev verir ve sonucu izler.'],
 ['Yetki, bilgi veya son durum belirsizliği','İlgili işlem bekler; gereken tek bilgi ya da karar doğru sahibinden istenir.']],[2.7,4.4],10.5)
h('Hamle mailde varsa aksiyon var',2)
p('Yetki, yeni mesajın gerçek To veya CC alanında hamle@gamechangers.com.tr bulunmasına dayanır. Eski zincirde bulunmak, BCC, iletilmiş metindeki adres veya Ahmet’in kutusunda görünmek yeterli değildir. Başlangıçta ve her yanıt, üretim, revizyon, plan değişikliği ve Basecamp işleminden önce kontrol gerekir.')
h('Çözülmesi gereken görünürlük sınırı',2)
p('Hamle sonraki mailden çıkarılırsa o mail Hamle’ye hiç gelmeyebilir. Dolayısıyla yalnız Hamle’nin kutusuna bakarak zincirin son alıcılarını kesin bildiğimizi söyleyemeyiz. Eski görünen maili hâlâ geçerli yetki saymak senin kuralını karşılamaz.',bold=True)
p('Canlı öncesi, yetki geri çekme ve insan devralma Hamle’nin görebildiği açık bir durdurma olayıyla işletilmelidir. Son alıcı durumunu gösterecek başka bir kaynak ancak ayrıca yetkilendirilir ve doğrulanır. Bu çözüm kurulana kadar görünmeyen değişime bağlı ertelenmiş müşteri işlemleri dış etki üretmez. Hamle kendini otomatik CC’ye eklemez; genel yönlendirme veya Ahmet’in kutusunu gizli canlı kaynak olarak kullanma çözüm sayılmaz.')

page('A Yönetim ve sorumluluklar')
p('Amaç, ajanın hangi kararı verebildiğini ve hangi noktada kime devredeceğini uygulamada kesinleştirmektir. Kapsam ve rol kuralları hazırdır; güncel marka matrisi, harcama sınırları ve kabul sahipleri tamamlanmalıdır.')
phase_rows('A')
h('Sorumluların anlamı',2)
p('Excel’de Teknik Kurulum, bu oturumda hazırlanan işlerin uygulanmasını ve testini yürütecek teknik sorumluyu ifade eder; ayrı bir ekip atanmış olduğu anlamına gelmez. Tasarım sorumlusu ve işletim yedeği için kişi atanmalıdır. Ajan adının bir işte bulunması, o ajanın bugün çalıştığını göstermez.')
h('Bitmiş sayılma koşulu',2)
p('Her marka için talep verebilen kişi, onaylayan kişi, kapsam, Basecamp proje alanı ve istisna sahibi kayıtlıdır. Güncel bütçe ve canlı kapsamı belirlenmiştir. Bu tanımlar yalnız dokümanda kalmayıp uygulamanın karar kurallarına taşınmıştır.')
p('Öncelikler: P0 canlı öncesi zorunlu bağımlılık; P1 ilgili hizmetin kabulü için gerekli iş; P2 kapsam genişletme veya sonraki iyileştirme. P0 olmak bütün işlerin aynı anda yapılacağı anlamına gelmez.',small=True)

page('B Gmail bağlantısı ve kalan kontroller')
p('Ahmet’in kurumsal hesabı Cloud yönetiminde; Hamle canlı mail kimliğinde kalır. Desktop istemcisi ve Hamle OAuth izni kuruldu. Gerçek API okuma, sabit iç gönderim ve Ahmet’te teslim doğrulandı; sırlar kullanıcıya özel güvenli depoda saklandı.')
phase_rows('B',['B01','B02','B03','B04','B05','B06','B07','B11','B12'])
p('Yerel Desktop istemcisi ilk kontrollü bağlantı içindir. Bulutta çalışacak sürümün dönüş adresi, gizli bilgi saklama ve gerekiyorsa yeniden yetkilendirme işi D10’da ayrı izlenir. Şifre sohbete alınmaz; giriş ve yetki Google’ın kendi ekranında tamamlanır.',small=True)
p('Gmail yenilemesi ve ayrı süreçte okuma geçti. İzin iptali ile Basecamp yenileme/yeniden başlatma açık olduğundan B07 tamamlanmadı. Yalnız kendi giden mesajının engellenmesi canlı sınandı; gelen posta, ekler ve yetki geri çekme için E03/E04 açık.',small=True)

page('B Basecamp ve marka eşleştirmesi')
p('Hamle’nin özel iç projesinde API görev, atama, R1 revizyon ve dosya yükleme döngüsü tamamlandı; dosyanın tam içeriği ve görünürlük yeniden doğrulandı. Bu yalnız iç test alanının kanıtıdır. Marka projeleri ayrıca doğrulanır; geniş sağlayıcı izni genel müşteri yetkisi sayılmaz.')
phase_rows('B',['B08','B09','B10'])
h('Pilot kapsamındaki tüm markalar',2)
table(['Marka','İlk bilgi durumu','Canlıya açılmadan gerekli'],[
 ['WAT Mobilite','Ek öğrenmede 27 ayrıntılı konu.','Yetkili kişiler, güncel kapsam, proje, onaylı varlıklar.'],
 ['Rabam','Ek öğrenmede 26 ayrıntılı konu.','Yetkili kişiler, güncel kapsam, proje, onaylı varlıklar.'],
 ['ya da','Eski birleşik gruptaki 7 konu için marka eşlemesi açık.','Kaynak bazında ayrım, yetki ve paylaşılabilir bilgi sınırı.'],
 ['Multicosmetics','Ayrı marka kaydı; güncel girdiler eksik.','ya da ile ortak kişi/izin varsaymadan ayrı yetki.'],
 ['Altınmarka','Düşük mail hacmi kullanıcı tarafından belirtildi.','altinmarkacc.com ve güncel marka kayıtları.'],
 ['Bikaldi','Düşük mail hacmi kullanıcı tarafından belirtildi.','altinmarkacg.com ve güncel marka kayıtları.'],
 ['HAN Spaces','Düşük mail hacmi kullanıcı tarafından belirtildi.','Güncel yetki, hizmet ve proje dosyası.'],
 ['Gamechangers','İç marka ve ilk gerçek demo alanı.','İç brief, değerlendiren kişi ve kabul edilen çıktı türleri.']],[1.35,2.7,3.05],9.7)
p('Yedi dış marka ve Gamechangers iç alanı pilot kapsamındadır. Hazırlık ayrı izlenir; eksik marka hazır markayı durdurmaz. Kapsama dahil olmak müşteri işlemi yetkisi vermez.',small=True)

page('C Marka bilgisi ve kurumsal hafıza')
p('Geçmiş yazışmadan çıkarılan bilgi, bugünkü sözleşme ve onaylı marka kılavuzuyla karşılaştırılacaktır. Eski tercih güncel talimat gibi sunulmaz. Her bilgiye kaynak, tarih, marka, geçerlilik ve onay durumu eklenir.')
phase_rows('C')
h('Marka dosyasının asgari içeriği',2)
p('Yetkili kişiler ve roller; hizmet kapsamı; ton ve yasak ifadeler; doğrulanmış ürün iddiaları; logo, font ve görsel hakları; kanal ölçüleri; son onaylı örnekler; açık işler; özel günler ve iletişim tercihleri. Doğum günü gibi bilgiler yalnız yetkili kaynaktan alınır ve ihtiyaç kadar tutulur.')
p('Arşiv öğrenimi yeni müşteri işi başlatmaz. Hangi sayfanın, dizinin ve ekin işlendiği kaydedilir; okunamayan ekler listede kalır. Bilgi değişince eski sürümün geçerliliği sonlandırılır; kaynağa geri dönülebilir.',small=True)

page('D Bulutta kalıcı çalışma')
p('Hedef, bilgisayardan ve bu sohbetten bağımsız çalışabilen bir hizmettir. İlk teknik tercih Cloud Run üzerinde işleyici; kalıcı kayıt ve dosya deposu; zamanlayıcı ve sır kasasıdır. Ücretli kaynaklar, güncel yetki ve bütçeden sonra açılır.')
phase_rows('D')
p('Günde 21 kontrol: Türkiye saatiyle hafta içi 08.30–18.30, iki uç dahil her 30 dakikada bir. Mesai dışı talepler sonraki çalışma başlangıcında alınır. 18.30 turu kontrollü tamamlanır. Tatil davranışı ajans takvimine bağlanır; takvim yoksa hafta içi varsayımı açıkça kayıt edilir.',small=True)

page('E Hamle iş akışı ve takip')
p('Hamle’nin görevi yalnız mail cevaplamak değildir. İşin sahibi, bir sonraki adımı, bekleyen kararı, teslim sürümünü ve takip zamanını yönetir. Bir işi iki kez açmayı, eski sürümü göndermeyi ve insan devraldıktan sonra devam etmeyi engelleyen kayıt gerekir.')
phase_rows('E')
h('İş durumları',2)
p('Yeni talep → Yetki ve bilgi kontrolü → Brief hazır → Üretimde → İç kontrolde → Müşteri değerlendirmesinde → Revizyonda → Onaylandı → Teslim edildi → Kapandı. Bekleyen bilgi, kapsam dışı ve insan devraldı durumları ayrıca görünürdür.')
p('Her devirde iş kimliği, marka, kaynak mesaj, beklenen çıktı, son sürüm, teslim zamanı ve bir sonraki sorumlu taşınır. Gönderim sonucu belirsizse kör tekrar yapılmaz; önce sağlayıcı kaydı araştırılır.',small=True)

page('F Metin ve G grafik üretimi')
p('İlk canlı kapsam önerisi: tanımlı rutin müşteri iletişimi, onaylı kaynaklardan metin üretimi ve kabul edilmiş şablonlarla basit grafikler. Yeni kampanya stratejisi, iddiası doğrulanmamış ürün metni ve üretim aracı desteklemeyen işler istisna yoluna gider.')
h('Metin yazarı',2)
phase_rows('F')
h('Grafik tasarımcı',2)
phase_rows('G')
p('Metin ve grafik için “ürettim” beyanı yeterli değildir. Gerçek çıktı açılmalı, briefle karşılaştırılmalı, son dosya ve revizyon kaydı doğrulanmalıdır. Son kullanıcıya giden dosya ile Basecamp’teki onaylı sürüm aynı olmalıdır.',small=True)

page('H Animasyon ve çıktı kalitesi')
p('Animasyon ilk üç rolün sınırlı açılışından ayrı kabul edilir. Desteklenen iş türü ve araç seçimi, gerçek örnek üretimle netleşir. Grafik rolünün başarısı hareketli içerik kalitesini doğrulamaz.')
phase_rows('H')
h('Teslim öncesi kalite kontrolü',2)
table(['Çıktı','Geçmesi gereken kontroller'],[
 ['Mail ve brief','Doğru marka ve alıcı; açık kapsam; doğrulanmış tarih ve taahhüt; gerçek son sürüm; doğru ek; güncel katılım.'],
 ['Metin','Kaynaklı iddia; marka tonu; yazım; uzunluk ve kanal biçimi; brief maddelerinin karşılanması.'],
 ['Grafik','Doğru logo ve font; hatasız Türkçe; okunurluk; boyut ve güvenli alan; lisans; dosyanın gerçekten açılması.'],
 ['Animasyon','Baştan sona izleme; doğru süre ve oran; ses/altyazı; görsel hatalar; haklar; son kare ve dışa aktarma.'],
 ['Revizyon','İstenen değişikliklerin tek tek karşılanması; önceki doğru bölümlerin korunması; sürüm ve onay eşleşmesi.']],[1.25,5.85],10.5)
p('Kabul örnekleri her marka ve iş türü için ayrı tutulur. Tasarım ve motion değerlendirmesini yapacak insan atanmadıysa ilgili üretim türünün müşteri kabulü açık kalır. Otomatik kontroller insan kalite kararının verildiğini varsayamaz.')

page('I Test ve canlıya geçiş koşulları')
phase_rows('I')
h('Canlı açılış için zorunlu koşullar',2)
p('1. Gerçek Hamle Gmail ve gerekli Basecamp işlemleri, erişim yenileme dahil geçti.\n2. Güncel katılım, yetki geri çekme ve insan devralma uçtan uca çalışıyor.\n3. Marka, alıcı, proje, kaynak ve sürüm ayrımı doğrulandı.\n4. Açılacak hizmet gerçek çıktı ve revizyon testinden geçti.\n5. Bulutta kesinti, tekrar ve belirsiz gönderim güvenli biçimde yönetiliyor.\n6. Güncel bütçe ve fiilî harcama kontrolü var.\n7. Sorun bildirimi doğru insana ulaşıyor; durdurma ve geri alma planı çalışıyor.\n8. İlgili marka ve hizmetin iç ve gölge kanıtları var; açık kritik kusur yok.')
p('Yanlış alıcı veya marka, yetkisiz işlem, çift gönderim, gizli bilgi sızıntısı ve durdurma sonrası devam etme kritik kusurdur. Sabit kabul setinde sıfır kritik hata gerekir; bu, gelecekte hata olmayacağı garantisi değildir.',small=True)

page('Kabul testleri nasıl yapılacak')
p('Her testte kaynak, başlangıç koşulu, beklenen sonuç, gerçek sonuç, iş ve dosya kimliği, değerlendirici ve tarih kaydedilir. Kritik testler gerçek bağlantılı iç ortamda yapılır. Eski müşteri talebi yeni işlem yetkisi sayılmaz; gölge denemeleri dış etki üretmez.')
table(['Vaka','Beklenen sonuç'],[
 ['Hamle To veya CC içinde','Doğru marka ve yetkiyle yalnız bir iş açılır.'],
 ['Hamle yok veya yalnız BCC ve alıntıda','Yanıt, iş, revizyon veya Basecamp işlemi oluşmaz.'],
 ['Açık iş sırasında geri çekme','Görülebilen stop olayı sonraki işlemi durdurur; görünmeyen değişim çözümü ayrıca doğrulanır.'],
 ['İnsan işi devraldı','Ajanın bekleyen yazmaları iptal edilir; insanın yanıtı veya sürümü üzerine işlem yapılmaz.'],
 ['Teklif, şikâyet ve sosyal plan','Sırasıyla Ahmet; Cansu ve Ahmet; Özge ile çalışma yoluna gider. Karma talepler ayrıştırılır.'],
 ['Yanlış marka, kişi veya eski kaynak','Erişim ve alıcı kontrolü işlemi durdurur; doğru kaynağa veya sorumluya döner.'],
 ['Mail veya ekte talimat saldırısı','İçeriğin kuralları, yetkileri veya alıcıları değiştirme talebi uygulanmaz; içerik iş verisi olarak ele alınır.'],
 ['Aynı olay iki kez veya eşzamanlı geldi','Tek iş ve tek dış işlem oluşur; kayıt sahipliği doğrulanır.'],
 ['Gönderim sonrası bağlantı koptu','Kör tekrar yapılmaz; sağlayıcı kaydıyla sonuç uzlaştırılır.'],
 ['İzin iptali, kesinti ve yeniden başlama','Güvenli durur, bildirir ve son güvenilir kayıttan kontrollü sürer.'],
 ['Yanlış sürüm onayı ve iki revizyon','Onay doğru dosyaya bağlıdır; istenen değişiklikler izlenir, son sürüm teslim edilir.'],
 ['Yanlış ürün iddiası ve bozuk görsel','Kaynak veya görsel kalite kontrolü reddeder; düzeltilen gerçek dosya yeniden kontrol edilir.'],
 ['Bütçe sınırı ve planlı tur eksikliği','Yeni ücretli iş sınırda durur; atlanan tur sorumluya bildirilir ve telafi kaydı tutulur.']],[2.25,4.85],10.1)
p('Rutin kalite için önerilen ölçü: zorunlu brief maddelerinin tamamı karşılanmış, kaynaksız iddia yok, dosya açılabiliyor ve değerlendirici kullanıma uygun buluyor. Sayısal başarı oranı verilecekse vaka sayısı, iş türü ve hata tanımı birlikte yazılır.',small=True)

page('Hızlandırılmış çalışma takvimi')
p('Aşağıdaki tarihler kapasite ve hazır girdilere bağlı hedeflerdir; taahhüt değildir. Bağımlı iş geçmezse sonraki adım kayar. 9 Ekim başlangıcı artık bir durum fotoğrafıdır; tamamlanmayan işler ilk uygun çalışma oturumuna taşınır.')
table(['Tarih','Hedef çıktı','İlerlemenin koşulu'],[
 ['9 Ekim Cuma\nBaşlangıç','Gmail ve Basecamp API iç testleri geçti; bütçe ve Astra + max kararı alındı.','Bulut, faturalama, model hesabı ve marka girdileri tamamlanmalı.'],
 ['10 Ekim Cumartesi\nİç demo','Ahmet’ten gerçek iç talep → Hamle → metin/grafik → Basecamp → revizyon → teslim.','Gerçek araçlar bağlı; hedef bulut ortamı hazır. Yerel demo bulut kabulü sayılmaz.'],
 ['11 Ekim Pazar\nGölge ve istisna','Eski gerçek talepler üzerinde dış etki olmadan karşılaştırma; kritik hata ve toparlanma testleri.','Kaynaklar, değerlendiriciler ve kabul örnekleri hazır; kritik açık yok.'],
 ['12 Ekim Pazartesi\nKoşullu sınırlı canlı','Koşulları geçen markalarda Hamle, metin ve basit grafik rutinleri.','Sekiz canlı koşulu geçti; açılacak marka ve iş türü kayıtlı.'],
 ['İlk canlı iş günleri\nİstikrar','Kaçan talep, hata, tekrar, süre ve maliyet gözlemi; kapsamı koruyarak düzeltme.','En az bir tam planlı iş günü izlenir; hatalı iş türü durdurulur.'],
 ['Sonraki kapsam\nTarihi açık','Daha karmaşık grafik, motion ve diğer üretim türleri.','Her yeni hizmet kendi araç, kalite ve revizyon kabulünü geçer.']],[1.37,2.95,2.78],10.1)
h('Kritik bağımlılık sırası',2)
p('Gmail temel bağlantısı ve bütçe kararı tamamlandı. Sırada bütçenin uygulanması, çalışma ortamı ve araçlar → Hamle iş akışı → gerçek metin ve grafik → iç demo → gölge kabul → kapsam bazlı açılış var. Marka dosyaları ve Basecamp eşlemleri paralel tamamlanır.')
p('Hafta sonu testleri, hafta içi 21 planlı turun bir tam iş günü boyunca sorunsuz çalıştığını kanıtlamaz. Pazartesi gözlemi bu nedenle canlı kapsamı büyütmenin ayrı koşuludur.',small=True)

page('Senden ve ekipten beklenen girdiler')
p('Excel’in Girdiler ve Kararlar sayfasında yanıt alanları hazırdır. Şifre veya erişim anahtarı bu tabloya yazılmaz. Atanmış kişi ve mevcut kararlar korunur; bilinmeyen alanlar tahminle doldurulmaz.')
table(['Girdi','Kimden','Neyi açar'],[
 ['Astra + max için yetkili API hesabı ve güvenli erişim; gerçek çağrı denemesi.','Ahmet ve teknik sorumlu','Modelin hesapta çalıştığının doğrulanması.'],
 ['Onaylı ortak 50 USD sınırının uygulanması; faturalama sahibi ve sınırda durdurma.','Ahmet ve teknik sorumlu','Ücretli araçlar ve Cloud işletimi.'],
 ['WAT, Rabam, ya da, Multicosmetics, Altınmarka, Bikaldi, HAN Spaces ve iç marka için kişiler ve yetkiler.','Ahmet ve ekip','Marka bazlı talep/onay ayrımı.'],
 ['Her marka için doğru Basecamp projesi ve iç/müşteri görünür alanlar.','Ahmet veya proje yöneticileri','Doğru yere görev, yorum ve dosya.'],
 ['Kılavuzlar, onaylı işler, logolar, fontlar, varlıklar, hizmet kapsamı ve geçerli iddialar.','Marka ve tasarım sorumluları','Metin ve görsel doğruluğu.'],
 ['Cansu, Özge, tasarım değerlendiricisi ve teknik işletim yedeğinin açık sorumlulukları.','Ahmet ve ilgili kişiler','İstisna, kalite ve kesinti yönetimi.'],
 ['Görünmeyen CC çıkarma için açık durdurma/geri çekme süreci.','Ahmet ve teknik sorumlu','Ertelenmiş müşteri işlerinin yetki kontrolü.'],
 ['Doğum günleri, özel günler, ajans tatilleri ve bilgi kullanım tercihleri.','Yetkili ekip ve marka sahipleri','Proaktif takvim ve sabah bilgilendirmesi.'],
 ['İlk iç brief, kabul edilecek metin/grafik türleri ve gölge örnekleri.','Ahmet, Özge, tasarım sorumlusu','Gerçek demo ve kalite karşılaştırması.']],[3.78,1.35,1.97],10)
p('Yeni onay: 9–12 Ekim 2026, Pazartesi dahil Google Cloud, OpenAI ve görsel üretim için ortak toplam 50 USD. Otomatik yenilenmez. Faturalama, gerçek tüketim ve uygulamadaki durdurma henüz doğrulanmadı; karar alınması D02’yi tamamlamaz.',small=True)

page('Canlı işletim ve sorun yönetimi')
h('Her çalışma günü',2)
p('08.30 ilk turda yeni talepler, bekleyen işler, yaklaşan teslimler ve açık kararlar kontrol edilir. Her yarım saatte yeni olaylar ve takipler işlenir. Mevcut başlangıç seçimi, ilk başarılı 08.30 turundan sonra Hamle hesabından Ahmet’e sabah özetidir; henüz zamanlanmış değildir. 18.30 son turu başlatılır ve kontrollü tamamlanır.')
p('Sabah özeti Hamle hesabından gönderilir: bugün teslim edilecekler, gecikenler, ekipten beklenenler, yaklaşan özel günler ve yalnız karar gerektiren istisnalar. Müşteriye özel gün maili atma yetkisi ayrıca tanımlıdır; takvime bilgi eklemek otomatik gönderim yetkisi değildir.')
h('Gözetim göstergeleri',2)
table(['Gösterge','Tanım ve işlem'],[
 ['Kontrol kapsaması','Gerçekleşen / beklenen planlı tur. Normal tam iş gününde payda 21; atlanan tur teknik sorumluya bildirilir.'],
 ['Talep bütünlüğü','Yetkili gelen mesajlar ile iş kaydı veya açık istisna kaydı eşleştirilir; sahipsiz mesaj açık hata sayılır.'],
 ['Kritik güvenlik olayı','Yanlış alıcı/marka, yetkisiz işlem, tekrar veya durdurma ihlali; etkilenen akış hemen durdurulur.'],
 ['Zaman ve kalite','İlk yanıt, teslim ve revizyon süreleri; ilk değerlendirmede kabul; nedenleriyle düzeltmeler. Hedefler hizmete göre belirlenir.'],
 ['Maliyet','Model, görsel/video ve Cloud kullanımı birlikte izlenir; işlem öncesi limit ve gecikmeli fatura farkı kontrol edilir.']],[1.55,5.55],10.3)
h('Durdurma ve dönüş',2)
p('Sorunda ilgili iş türü veya marka durdurulur; sıradaki dış işlemler iptal edilir ve insan devralır. Son güvenilir iş kaydı ile sağlayıcı kayıtları karşılaştırılır. Gönderilmiş işlemi geri alınmış saymadan olay özeti çıkarılır. Düzeltme ve tekrar testten sonra kontrollü devam edilir.')
p('Bulut bütçe alarmı tek başına harcamayı durdurmayabilir. Seçilen hizmetlerde uygulanabilir harcama tavanı, kota ve uygulama sınırları ayrıca doğrulanır; bütün sağlayıcılar tek toplam bütçede izlenir. Hamle’nin AI rolü açık kalır; yapılmamış insan görüşmesi veya kontrol yapılmış gibi anlatılmaz.',small=True)

page('Yeni Mac ve yeni konuşmaya devir')
p('Bu belge 10 Ekim’de yeni bilgisayara devir için güncellendi. İlerleme, 9 Ekim 23.59’a kadar doğrulanan işlemleri gösterir. Word ve Excel bilgi ve iş takibini taşır; yazılımı, erişim sırlarını veya hesap bağlantılarını taşımaz.',bold=True)
table(['Kalem','Mevcut durum ve yeni cihaz adımı'],[
 ['Kalıcı uzak kayıtlar','Cloud projesi gamechangers-hamle ve OAuth uygulaması Gamechangers Hamle mevcut. Basecamp iç proje 49184053, Hamle kişi 53358100. Aynı kayıtlar doğrulanır; cihaz değişti diye yeniden oluşturulmaz.'],
 ['Bu Mac’teki erişimler','Gmail kimliği macOS Keychain’de, Basecamp bağlantısı yerel keyring’de. Yeni Mac’e taşındığı varsayılmaz. Doğru Hamle kimliği, özel saklama, okuma/yazma ve yenileme yeni cihazda yeniden doğrulanır.'],
 ['OpenAI hesabı','OpenAI Developers kuruldu ve proje seçim akışı açıldı. Hedef seçilmedi, API anahtarı oluşturulmadı. Doğru kuruluş/proje ve güvenli anahtar kurulumu yeni Mac’te tamamlanır; model erişimi henüz doğrulanmadı.'],
 ['Onaylı kararlar','Astra + max denemesi onaylı. Google Cloud, OpenAI ve görsel üretim için 12 Ekim Pazartesi dahil ortak toplam 50 USD. Gerçek tüketim ve uygulama sınırı henüz kurulmadı; otomatik yenileme yok.'],
 ['Çalışma dosyaları','work/project-tracker, work/hamle-workflow, work/hamle-gmail-connector ve outputs/gamechangers-pilot dosyaları eski Mac’te. Gereken paket güvenli biçimde aktarılmalı; sırlar ve ham özel mailler sohbet ekine konmamalı.'],
 ['Yeniden doğrulama','54 yeni posta okuma/normalleştirme testi geçti. Canlıda yalnız kendi giden mesajını engelleme sınandı; gelen posta, ekler ve görünmeyen CC çıkarma çözümü açık. Yeni ortamda testler tekrarlanır.']],[1.25,5.85],10.2)
h('Yeni konuşmanın ilk çalışma sırası',2)
p('1. İki belgeyi oku; 27 puan, 19 tamam iş ve 0 canlı ajan bilgisini esas al.\n2. Eski uzak proje ve iç test kimliklerini doğrula; yerel bağlantıları güvenle yeniden kur.\n3. Model hesabını ve ortak bütçe kontrolünü tamamla.\n4. E03/E04 yetki kontrollü posta akışı ile Cloud işletimini kur; marka girdilerini topla.\n5. İç demo ve gölge kabulü geçmeden müşteri akışını açma.')
p('Yeni bir ChatGPT konuşması açılması, eski çalışma ortamını veya çalışan bir bulut hizmetini oluşturmaz. Hamle yalnız güncel gerçek To/CC katılımında iş yapar; yetki geri çekme ve insan devralma kontrolleri her dış işlem öncesinde geçmelidir.',small=True)

page('Takibi güncelleme ve kanıt dizini')
h('Excel nasıl kullanılacak',2)
p('Özet sayfası 100 puanlık ilerlemeyi ve canlı ajan sayısını ayrı gösterir. İş Listesi sayfasında durum, sorumlu, hedef tarih ve sonraki adım güncellenir. Bir işin kabul koşulu ve kanıtı tamamlanınca Tamamlandı alanı 1 yapılır; yüzde otomatik hesaplanır. Girdiler ve Kararlar sayfasında kararlar, Yeni Mac Devir sayfasında taşınacak kayıtlar ve cihaz kontrolleri tutulur.')
p('Hedef tarihi boş bırakılmış işler henüz takvim taahhüdü değildir. Bu rapordaki hafta sonu hedefi otomatik son tarih olarak bütün satırlara yazılmadı. Yeni iş eklenirse ağırlıklar, özet formül aralıkları ve belge sürümü birlikte güncellenir. Durum metni ile 0/1 alanı çelişirse önce kanıt kontrol edilir.')
p('Bağımlılıklar açılacak marka ve hizmet için uygulanır. A04, B10 ve C05 gibi portföy işleri ancak tüm kapsam bitince 1 olur; ilk sınırlı açılış için ilgili markanın sekiz canlı koşulu IN13 alanında ayrı kabul kaydıyla gösterilir. Böylece eksik bir marka hazır markayı bekletmez.',small=True)
h('Durum kanıtları',2)
p('Aşağıdaki dosyalar çalışma alanındaki kayıtları işaret eder. Excel’de ilgili iş satırına da kaynak bağlanmıştır. Hesap erişim anahtarları ve ham müşteri yazışmaları bu rapora eklenmemiştir.',small=True)
sources=[
 ('Gmail API, teslim ve yenileme','outputs/Hamle_Gmail_API_Kaniti_2026-10-09.json; outputs/Hamle_Gmail_Refresh_Kaniti_2026-10-09.json'),
 ('Basecamp görev dosya ve revizyon kabulü','outputs/Hamle_Basecamp_B09_API_Kaniti_2026-10-09.json'),
 ('Güncel bütçe ve model kararı','outputs/Hamle_Butce_Onayi_2026-10-09.json; outputs/Hamle_Model_API_On_Kontrol_2026-10-09.json'),
 ('Güncel Cloud ve OAuth durumu','outputs/Hamle_Google_Cloud_Kurulum_2026-10-09.json'),
 ('Kalıcı Basecamp API testi','outputs/Hamle_Kalici_Baglanti_Hazirlik_2026-10-09.json'),
 ('Gmail ve Basecamp iç iletişim testleri','outputs/Hamle_Baglanti_Durumu_2026-10-09.json'),
 ('Öğrenme kapsamı ve sınırları','outputs/gamechangers-pilot/coverage.json'),
 ('Operasyon kural testleri','outputs/gamechangers-pilot/offline/test-results.json'),
 ('Gmail yardımcı testleri','work/hamle-gmail-connector/test-results.json'),
 ('Dört rol ve politika dosyaları','outputs/gamechangers-pilot/roles'),
 ('Bağımsız durum denetimi','work/project-tracker/status-audit.json')]
for label,path in sources: p(label+' — '+path,small=True)
h('Uygulama için resmî başvuru kaynakları',2)
for label,url in [
 ('Google Gmail bağlantı başlangıcı','https://developers.google.com/workspace/gmail/api/quickstart/python'),
 ('Google OAuth kullanıcı yetkisi ve yenileme','https://developers.google.com/identity/protocols/oauth2'),
 ('Cloud Run çalışma ortamı','https://docs.cloud.google.com/run/docs/overview/what-is-cloud-run'),
 ('Cloud Run zamanlanmış işler','https://docs.cloud.google.com/run/docs/execute/jobs-on-schedule'),
 ('Cloud sır yönetimi','https://docs.cloud.google.com/run/docs/configuring/services/secrets'),
 ('Cloud bütçe ve uyarı seçenekleri','https://docs.cloud.google.com/billing/docs/how-to/budgets')]: link(label,url)
p('Resmî kaynaklar hedef yaklaşımı açıklar; hesap kurulumunun kanıtı test kayıtlarıdır. Bu dosya yenilemesi ek hesap işlemi veya ücretli kaynak oluşturmaz.',small=True)

assert total==100 and pts==D['currentSummary']['completedWeight']
doc.save(OUT/'Hamle_Proje_Plani_2026-10-10_Yeni_Mac_Devir.docx')
print(str(OUT/'Hamle_Proje_Plani_2026-10-10_Yeni_Mac_Devir.docx'))
